"""Generar informe usando las métricas guardadas, sin inventar resultados."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def main():
    metrics=json.loads((ROOT/'resultados/metricas.json').read_text())
    audit=json.loads((ROOT/'resultados/auditoria_dataset.json').read_text())
    doc=Document()
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.9)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[name]; style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Title'].font.size=Pt(25)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.add_heading('Clasificación de imágenes de fuego mediante una CNN en Android',0)
    doc.add_paragraph('Bosque Alerta',style='Subtitle')
    doc.add_paragraph('Proyecto final de Aprendizaje Profundo\nUniversidad: [Completar]\nIX semestre')
    doc.add_paragraph('Integrantes\n1. [Nombre y apellidos]\n2. [Nombre y apellidos]\n3. [Nombre y apellidos]\n\nDocente: [Completar]\nFecha de presentación: [Completar]')
    doc.add_heading('Resumen',1)
    doc.add_paragraph(f'Se desarrolló un prototipo Android que clasifica fotografías en fuego y sin fuego mediante una red neuronal convolucional entrenada desde cero. Se utilizó FIRE Dataset de Ahmed Saied, con 999 imágenes originales. Tras la auditoría quedaron {audit["valid_unique"]} imágenes únicas válidas. La evaluación reservada obtuvo una exactitud de {metrics["accuracy"]:.2%} y un recall de fuego de {metrics["fuego"]["recall"]:.2%}. El modelo se convirtió a TensorFlow Lite para inferencia local. Estos resultados describen la partición del dataset utilizada y no demuestran rendimiento operativo en incendios forestales de Puno.')
    doc.add_page_break()
    doc.add_heading('1 Planteamiento del problema',1)
    doc.add_paragraph('La revisión visual de fotografías puede apoyar la identificación inicial de señales de fuego. El proyecto estudia si una CNN pequeña permite clasificar una fotografía y ejecutar esa clasificación en un teléfono sin conexión. Su aplicación propuesta son escenas de posibles incendios forestales; las etiquetas disponibles solamente distinguen presencia y ausencia de fuego. Por ello no se afirma que el sistema determine el tipo de incendio, su ubicación o su gravedad.')
    doc.add_heading('2 Objetivos',1)
    doc.add_paragraph('El objetivo general es implementar y evaluar una aplicación Android de clasificación binaria de imágenes mediante una CNN. Los objetivos específicos son organizar y auditar el dataset, construir particiones reproducibles, entrenar una arquitectura convolucional, medir su desempeño en prueba, exportarla a TensorFlow Lite y empaquetarla en un APK.')
    doc.add_heading('3 Fundamento del método',1)
    doc.add_paragraph('Una convolución aplica filtros aprendidos sobre regiones locales de la imagen. La activación ReLU introduce no linealidad y el pooling reduce la resolución espacial. Las capas densas combinan las características para generar una salida sigmoide entre cero y uno. Se optimiza la entropía cruzada binaria con Adam. La salida se interpreta como puntaje de fuego: un valor mayor o igual a 0.5 produce la clase fuego. No se realizó calibración probabilística.')
    doc.add_heading('4 Datos y procedencia',1)
    doc.add_paragraph('FIRE Dataset fue publicado por Ahmed Saied en Kaggle con la referencia phylake1337/fire-dataset, versión 1. La ficha declara licencia CC0. La descarga se conserva en dataset/fire-dataset-v1.zip y procedencia.json registra la URL y SHA-256 del archivo. Se encontraron 755 imágenes con fuego y 244 sin fuego. La unidad de clasificación es una imagen completa; no se dispone de cajas delimitadoras ni máscaras.')
    doc.add_paragraph(f'La auditoría verificó decodificación, convirtió imágenes a RGB y comparó hashes de píxeles y dimensiones. Excluyó {len(audit["duplicates_excluded"])} duplicados exactos y {len(audit["invalid"])} archivos inválidos. No se ha certificado la ausencia de imágenes casi duplicadas o de escenas relacionadas. Una evaluación más estricta debe agrupar por escena, evento o secuencia antes de dividir.')
    doc.add_heading('5 Diseño experimental',1)
    doc.add_paragraph('Se utilizó semilla 42 y una división estratificada aproximada de 70 % entrenamiento, 15 % validación y 15 % prueba. El manifiesto particiones.csv permite identificar cada imagen y comprobar la exclusión entre subconjuntos. Solo entrenamiento recibe aumentos aleatorios y ponderación inversa de frecuencia de clases. El conjunto de validación controla la parada temprana; el de prueba se utiliza después del entrenamiento, sin ajustar el umbral a sus resultados.')
    table(doc,['Subconjunto','Imágenes'],[[k,str(v)] for k,v in metrics['split_counts'].items()])
    doc.add_paragraph('El preprocesamiento aplica orientación EXIF, conversión RGB, redimensionamiento bilineal a 96 × 96 píxeles y división por 255. El redimensionamiento estira la imagen; puede deformar detalles. Android usa el mismo tamaño y escala; las implementaciones de interpolación pueden presentar pequeñas diferencias. No se mide en este experimento el impacto de esas diferencias sobre todo el dataset.')
    doc.add_heading('6 Arquitectura y entrenamiento',1)
    table(doc,['Etapa','Configuración'],[
        ['Entrada','96 × 96 × 3, float32 RGB en [0,1]'],
        ['Aumento solo durante entrenamiento','Volteo horizontal y rotación aleatoria ±18 grados'],
        ['Bloque 1','Conv2D 16 filtros 3 × 3, ReLU, padding same; MaxPool 2 × 2'],
        ['Bloque 2','Conv2D 32 filtros 3 × 3, ReLU; MaxPool 2 × 2'],
        ['Bloque 3','Conv2D 64 filtros 3 × 3, ReLU; MaxPool 2 × 2'],
        ['Clasificador','Flatten; Dense 64 ReLU; Dropout 0.4; Dense 1 sigmoide'],
        ['Optimización','Adam, tasa 0.001; lote 32; entropía cruzada binaria'],
        ['Parada','Máximo 25 épocas; paciencia 5 en val_loss; restaura mejores pesos']])
    doc.add_paragraph(f'El modelo contiene {metrics["parameters"]:,} parámetros y ejecutó {metrics["epochs_executed"]} épocas. Se empleó TensorFlow {metrics["tensorflow"]}. Las versiones completas del entorno se conservan en requirements-lock.txt. La reproducibilidad numérica puede variar entre hardware y versiones, aunque se habilitaron operaciones deterministas y se registró la semilla.')
    doc.add_heading('7 Resultados de prueba',1)
    table(doc,['Indicador','Valor'],[
        ['Exactitud',f'{metrics["accuracy"]:.4f}'],
        ['Precisión fuego',f'{metrics["fuego"]["precision"]:.4f}'],
        ['Recall fuego',f'{metrics["fuego"]["recall"]:.4f}'],
        ['F1 fuego',f'{metrics["fuego"]["f1-score"]:.4f}'],
        ['F1 macro',f'{metrics["macro avg"]["f1-score"]:.4f}'],
        ['ROC AUC',f'{metrics["roc_auc"]:.4f}']])
    cm=metrics['confusion_matrix']
    table(doc,['Clase real','Predijo sin fuego','Predijo fuego'],[['Sin fuego',str(cm[0][0]),str(cm[0][1])],['Fuego',str(cm[1][0]),str(cm[1][1])]])
    doc.add_paragraph('La matriz usa filas como clases reales y columnas como predicciones. El recall de fuego mide qué fracción de imágenes con fuego detectó el modelo; los falsos negativos son imágenes con fuego clasificadas sin fuego. La precisión mide qué fracción de predicciones de fuego fue correcta. F1 macro pondera por igual las clases y ayuda a interpretar el desbalance. Las métricas corresponden a una sola partición, sin intervalos de confianza ni validación externa.')
    doc.add_picture(str(ROOT/'resultados/evaluacion.png'),width=Inches(6.2))
    doc.add_paragraph('Figura 1. Pérdida durante entrenamiento y validación y matriz de confusión de prueba. Fuente: ejecución del proyecto.',style='Caption')
    doc.add_heading('8 Conversión y aplicación Android',1)
    doc.add_paragraph(f'El archivo bosque_alerta.tflite recibe un tensor float32 de forma [1, 96, 96, 3] y devuelve un puntaje sigmoide. La comparación con Keras en prueba obtuvo una diferencia absoluta máxima de {metrics["tflite_max_absolute_error"]:.8f} y coincidencia de etiquetas de {metrics["tflite_prediction_agreement"]:.2%}. Esta verificación valida la conversión sobre el mismo tensor; no sustituye pruebas de extremo a extremo en el teléfono.')
    doc.add_paragraph('La aplicación está escrita en Java y utiliza el intérprete TensorFlow Lite 2.17.0. Requiere Android 9 o posterior. El usuario selecciona una foto con el selector del sistema; la app muestra imagen, clase y puntaje. El modelo se ejecuta en un hilo de trabajo y las imágenes permanecen en el dispositivo. El APK de depuración es para demostración académica. Las pruebas manuales de instalación, orientación, cancelación, archivos inválidos y funcionamiento en modo avión se registran en apk/PRUEBAS_MANUALES.md.')
    doc.add_heading('9 Limitaciones y mejoras',1)
    doc.add_paragraph('La muestra es pequeña, desbalanceada y no constituye un muestreo representativo de todas las condiciones forestales. Pueden existir sesgos de color, fuente y contexto, así como correlación entre imágenes. Humo sin llamas, atardeceres, luces y objetos anaranjados requieren evaluación específica. Se debe recopilar un conjunto externo de escenas forestales locales, revisar los casos erróneos, agrupar por evento y comparar con una CNN preentrenada antes de proponer uso fuera del aula.')
    doc.add_heading('10 Organización del equipo',1)
    table(doc,['Responsable propuesto','Contribución'],[['Integrante 1','Datos, auditoría y exploración'],['Integrante 2','CNN, entrenamiento y evaluación'],['Integrante 3','Android, pruebas de dispositivo y demo'],['Los tres','Revisión del informe y exposición']])
    doc.add_heading('11 Conclusiones',1)
    doc.add_paragraph(f'La implementación integra preparación de datos, entrenamiento de una CNN desde cero, evaluación y exportación para Android. La exactitud observada fue {metrics["accuracy"]:.2%} en la partición de prueba. La coincidencia entre Keras y TensorFlow Lite permite conservar el comportamiento del modelo para las entradas verificadas. El alcance demostrado es clasificación binaria sobre este dataset; la validación externa en escenas forestales y las pruebas en dispositivos son necesarias para ampliar esa conclusión.')
    doc.add_heading('12 Referencias',1)
    for line in ['Saied, A. FIRE Dataset. Kaggle, versión 1. https://www.kaggle.com/datasets/phylake1337/fire-dataset',
        'TensorFlow. tf.keras.Sequential. https://www.tensorflow.org/api_docs/python/tf/keras/Sequential',
        'TensorFlow. TensorFlow Lite inference. https://github.com/tensorflow/tensorflow/blob/master/tensorflow/lite/g3doc/guide/inference.md']:
        doc.add_paragraph(line)
    export_pdf_markdown(doc)
    print('Informe generado')

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,text in zip(t.rows[0].cells,headers): c.text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=text
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            props.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EBEF' if i==0 else ('F5F6F7' if i%2==0 else 'FFFFFF')); props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(6); p.paragraph_format.space_before=Pt(6)
                for run in p.runs: run.bold=i==0; run.font.size=Pt(10)
    doc.add_paragraph('')

def export_pdf_markdown(doc):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from docx.text.paragraph import Paragraph as WordParagraph
    from docx.table import Table as WordTable
    from xml.sax.saxutils import escape
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='BodySpanish',fontName='Helvetica',fontSize=10.5,leading=15,spaceAfter=8))
    styles['Title'].fontSize=24; styles['Title'].leading=30; styles['Title'].alignment=0
    styles['Heading1'].fontSize=15; styles['Heading1'].leading=19; styles['Heading1'].spaceBefore=14
    styles.add(ParagraphStyle(name='CellSpanish',fontName='Helvetica',fontSize=9,leading=12))
    story=[]; markdown=[]
    for element in doc.element.body:
        if element.tag==qn('w:p'):
            p=WordParagraph(element,doc)
            if element.xpath('.//w:br[@w:type="page"]'):
                story.append(PageBreak()); continue
            if element.xpath('.//w:drawing'):
                story.append(Image(str(ROOT/'resultados/evaluacion.png'),width=475,height=190))
                markdown.append('![Evaluación](../resultados/evaluacion.png)'); continue
            text=p.text.strip()
            if not text: continue
            name=p.style.name
            style=styles['Title'] if name=='Title' else styles['Heading1'] if name.startswith('Heading') else styles['BodySpanish']
            story.append(Paragraph(escape(text).replace('\n','<br/>'),style))
            markdown.append(('# ' if name=='Title' else '## ' if name.startswith('Heading') else '')+text)
        elif element.tag==qn('w:tbl'):
            t=WordTable(element,doc)
            data=[[Paragraph(escape(c.text),styles['CellSpanish']) for c in row.cells] for row in t.rows]
            widths=[475/len(data[0])]*len(data[0])
            pdf_table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E6EBEF')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.extend([pdf_table,Spacer(1,10)])
            rows=[[c.text.replace('\n',' ') for c in row.cells] for row in t.rows]
            markdown.append('\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(r)+' |' for r in rows[1:]]))
    target=ROOT/'informe/Informe_Bosque_Alerta.pdf'
    def footer(canvas,document):
        canvas.setFont('Helvetica',9); canvas.setFillColor(colors.HexColor('#666666'))
        canvas.drawRightString(A4[0]-60,32,str(document.page))
    SimpleDocTemplate(str(target),pagesize=A4,leftMargin=60,rightMargin=60,topMargin=50,bottomMargin=50,title='Clasificación de imágenes de fuego mediante una CNN en Android',author='Equipo Bosque Alerta').build(story,onFirstPage=footer,onLaterPages=footer)
    (ROOT/'informe/Informe_Bosque_Alerta.md').write_text('\n\n'.join(markdown),encoding='utf-8')

if __name__=='__main__': main()
